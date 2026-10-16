import os


def parse_pptx(file_path, cfg):
    """
    Parse PPTX files: extract slides, text, tables, notes.
    Returns a normalized document dict.
    """
    from pptx import Presentation

    prs = Presentation(file_path)
    result = {
        'metadata': {},
        'pages': [],
        'sections': [],
        'tables': [],
        'totalPages': len(prs.slides),
        'documentType': 'PPTX',
    }

    all_text = []
    slide_data = []
    extracted_tables = []

    for i, slide in enumerate(prs.slides):
        slide_texts = []
        title = ''

        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    text = paragraph.text.strip()
                    if text:
                        slide_texts.append(text)
                        all_text.append(text)
                        # Heuristic: first short text in a title-like shape is the title
                        if not title and shape.shape_type is not None and len(text) < 100:
                            if hasattr(shape, 'name') and ('title' in (shape.name or '').lower() or 'Title' in (shape.name or '')):
                                title = text

            if shape.has_table:
                try:
                    table = shape.table
                    headers = [cell.text.strip() for cell in table.rows[0].cells]
                    rows = []
                    for row in table.rows[1:]:
                        rows.append([cell.text.strip() for cell in row.cells])
                    extracted_tables.append({
                        'slide': i + 1,
                        'tableIndex': len(extracted_tables) + 1,
                        'headers': headers,
                        'rows': rows,
                        'rowCount': len(rows),
                        'colCount': len(headers),
                    })
                except Exception:
                    pass

        # If no title detected, use first text as title
        if not title and slide_texts:
            title = slide_texts[0]

        # Extract notes
        notes_text = ''
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes_text = slide.notes_slide.notes_text_frame.text.strip()

        slide_data.append({
            'slide': i + 1,
            'title': title,
            'content': slide_texts,
            'notes': notes_text,
        })

    result['pages'] = slide_data
    result['tables'] = extracted_tables
    result['fullText'] = '\n\n'.join(all_text)[:cfg['max_extracted_text_chars']]

    return result
